# -*- coding: utf-8 -*-
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import nba_home


class Dates(unittest.TestCase):
    def test_focus_before_results_is_rejected(self):
        with self.assertRaises(ValueError):
            nba_home.validate_dates("20261022", "20261021")

    def test_next_day_has_no_warning(self):
        results, focus, warning = nba_home.validate_dates("20261021", "20261022")
        self.assertEqual(results, "20261021")
        self.assertEqual(focus, "20261022")
        self.assertIsNone(warning)

    def test_bad_calendar_day(self):
        with self.assertRaises(ValueError):
            nba_home.parse_yyyymmdd("20260231")


class Commands(unittest.TestCase):
    def test_results_steps_receive_the_date(self):
        crawler = nba_home.default_command(1, "20261021")
        ox = nba_home.default_command(2, "20261021")
        teams = nba_home.default_command(3, "20261021")
        self.assertEqual(crawler[1:], ["執行範圍爬蟲.py", "20261021", "20261021"])
        self.assertEqual(ox[1:], ["批次轉換正負盤.py", "20261021~20261021"])
        self.assertEqual(teams[1:], ["批次轉換各隊正負.py", "20261021~20261021"])

    def test_focus_steps_do_not_receive_the_date(self):
        for step in (4, 5, 6):
            cmd = nba_home.default_command(step, "20261022")
            self.assertNotIn("20261022", cmd[1:])

    def test_menu_command_is_used_when_run_all_exists(self):
        class Menu:
            def _build_command(self, step, start, end):
                if step == 6:
                    raise ValueError("no step")
                return ["python", f"step{step}.py", start]

        self.assertEqual(
            nba_home.command_for(1, "20261021", Menu()),
            ["python", "step1.py", "20261021"],
        )
        self.assertIsNone(nba_home.command_for(6, "20261022", Menu()))

    def test_mlb_menu_is_recognized(self):
        class Mlb:
            STEP_NAMES = {4: "ESPN MLB Batting"}

        class Nba:
            STEP_NAMES = {4: "ESPN Player Stats/Injuries"}

        self.assertTrue(nba_home.is_mlb_menu(Mlb()))
        self.assertFalse(nba_home.is_mlb_menu(Nba()))


class Plan(unittest.TestCase):
    def test_failure_does_not_stop_later_steps(self):
        calls = []

        def runner(cmd):
            calls.append(cmd[1])
            return 1 if cmd[1] == "批次轉換正負盤.py" else 0

        log = nba_home.run_plan("20261021", "20261022", runner, run_all=None)
        self.assertEqual(
            [name for name, _status, _detail in ((item[0], item[1], item[2]) for item in log)],
            [1, 2, 3, 4, 5, 6],
        )
        self.assertIn("批次轉換正負盤.py", calls)
        self.assertIn("espn_player_stats.py", calls)
        self.assertEqual(log[1][1], "FAILED")
        self.assertEqual(log[2][1], "OK")

    def test_missing_step_is_skip_not_failure(self):
        class Menu:
            def _build_command(self, step, start, end):
                if step == 6:
                    raise ValueError("missing")
                return [sys.executable, f"s{step}.py"]

        log = nba_home.run_plan("20261021", "20261022", lambda cmd: 0, Menu())
        self.assertEqual(log[-1], (6, "SKIP", "選單沒有這一步"))

    def test_dry_run_main_does_not_launch_scripts(self):
        launched = []
        original = nba_home.subprocess_runner
        nba_home.subprocess_runner = lambda cmd: launched.append(cmd) or 0
        try:
            code = nba_home.main(["--dry-run", "20261021", "20261022"])
        finally:
            nba_home.subprocess_runner = original
        self.assertEqual(code, 0)
        self.assertEqual(launched, [])


class ArchiveLayout(unittest.TestCase):
    def _book(self):
        from openpyxl import Workbook
        return Workbook()

    def test_record_layout_plans_archive_and_date(self):
        import nba_archive
        wb = self._book()
        master = wb.active
        master.title = "總表"
        master["A1"] = "20261020"
        record = wb.create_sheet("紀錄")
        record["A3"] = "第1場"
        kept = wb.create_sheet("各隊total")
        kept["A2"] = "賽爾提克"
        layout = nba_archive.classify_workbook(wb)
        self.assertEqual(layout["kind"], "record")
        results = nba_archive.actions_for("record", "20261021", "results")
        focus = nba_archive.actions_for("record", "20261022", "focus")
        self.assertIn("紀錄 K1=20261021", results)
        self.assertIn("總表 A1=20261022", focus)
        self.assertNotIn("貼上值", " ".join(focus))
        self.assertNotIn("各隊total", " ".join(results + focus))

    def test_calendar_layout_only_recalculates(self):
        import nba_archive
        wb = self._book()
        master = wb.active
        master.title = "總表"
        master["A3"] = "天數"
        master["A4"] = "日期"
        master["A6"] = "一"
        layout = nba_archive.classify_workbook(wb)
        self.assertEqual(layout["kind"], "calendar")
        self.assertEqual(nba_archive.actions_for("calendar", "20261021", "results"), ["重算整本"])

    def test_record_with_label_in_a1_does_not_overwrite_a1(self):
        import nba_archive
        wb = self._book()
        master = wb.active
        master.title = "總表"
        master["A1"] = "NBA總表"
        record = wb.create_sheet("紀錄")
        record["A2"] = "第1場"
        layout = nba_archive.classify_workbook(wb)
        self.assertEqual(layout["kind"], "record_keep_a1")
        results = nba_archive.actions_for("record_keep_a1", "20261021", "results")
        self.assertIn("紀錄 K1=20261021", results)
        self.assertFalse(any(item.startswith("總表 A1=") for item in results))

    def test_unknown_layout_has_no_writes(self):
        import nba_archive
        wb = self._book()
        wb.active["A1"] = "其他表"
        layout = nba_archive.classify_workbook(wb)
        self.assertEqual(layout["kind"], "unknown")
        self.assertEqual(nba_archive.actions_for("unknown", "20261021", "results"), [])

    def test_refuses_mlb_workbook_path(self):
        import nba_archive
        os.environ["NBA_STATS_FILE"] = r"C:\Users\curti\OneDrive\MLB26\MLB26-27數據.xlsx"
        try:
            with self.assertRaises(SystemExit):
                nba_archive.excel_path()
        finally:
            os.environ.pop("NBA_STATS_FILE", None)

    def test_missing_file_on_dry_run_is_skip(self):
        import nba_archive
        label, status, detail = nba_archive.run_archive(
            r"C:\missing\NBA26-27爬蟲數據.xlsx", "20261021", "results", True,
        )
        self.assertEqual(status, "SKIP")
        self.assertIn("找不到", detail)
        self.assertEqual(label, "歸檔賽果日")


if __name__ == "__main__":
    unittest.main()
