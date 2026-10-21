# -*- coding: utf-8 -*-
"""
爬取 ESPN 球員狀態，並一併更新球員傷兵。
- 球員狀態：Player Stats + Shooting，寫入「球員狀態」（一列一球員）
- 球員傷兵：寫入「球員傷兵」（球隊、球員、傷勢、狀態、日期）
- 球員狀態任一隊抓失敗就不覆寫這兩個分頁
"""

import os
import re
import sys

from bs4 import BeautifulSoup

from espn_common import (
    NBA_TEAMS,
    SEASON_LABEL,
    SEASON_TYPE,
    SEASON_TYPE_NAME,
    SEASON_YEAR,
    cell_value,
    excel_path,
    fetch_html,
    injuries_url,
    stats_url,
)

COUNTING_FIELDS = [
    "GP", "GS", "MIN", "PTS", "OR", "REB", "AST", "STL", "BLK", "TO", "PF", "AST/TO",
]
SHOOTING_FIELDS = [
    "FGM", "FGA", "FG%", "3PM", "3PA", "3P%", "FTM", "FTA", "FT%",
    "2PM", "2PA", "2P%", "SC-EFF", "SH-EFF",
]


def _text(cell):
    return " ".join(cell.get_text(" ", strip=True).split())


def _headers(table):
    return [_text(th).upper() for th in table.find_all("th")]


def _data_rows(table):
    body = table.find("tbody") or table
    rows = []
    for tr in body.find_all("tr"):
        cells = tr.find_all("td")
        if not cells:
            continue
        rows.append([_text(cell) for cell in cells])
    return rows


def _blocks(html):
    soup = BeautifulSoup(html, "html.parser")
    tables = soup.find_all("table")
    headers = [_headers(table) for table in tables]
    basic = shoot = None
    basic_need = {"GP", "PTS", "REB", "AST"}
    shoot_need = {"FGM", "FG%", "3PM", "FT%"}
    for index in range(len(tables) - 1):
        left, right = headers[index], headers[index + 1]
        if left == ["NAME"] and basic_need.issubset(set(right)) and basic is None:
            basic = (_data_rows(tables[index]), right, _data_rows(tables[index + 1]))
        if left == ["NAME"] and shoot_need.issubset(set(right)) and shoot is None:
            shoot = (_data_rows(tables[index]), right, _data_rows(tables[index + 1]))
    return basic, shoot


def _zip_map(headers, values):
    width = min(len(headers), len(values))
    return {headers[i]: values[i] for i in range(width)}


def parse_player_stats(html, team):
    basic, shoot = _blocks(html)
    if not basic or not shoot:
        return []
    names, basic_headers, basic_rows = basic
    shoot_names, shoot_headers, shoot_rows = shoot
    rows = []
    for index, name_cells in enumerate(names):
        if not name_cells:
            continue
        name = name_cells[0]
        if not name or name in ("Total", "NAME"):
            continue
        basic_map = _zip_map(basic_headers, basic_rows[index] if index < len(basic_rows) else [])
        shoot_map = {}
        shoot_index = next(
            (i for i, cells in enumerate(shoot_names) if cells and cells[0] == name),
            None,
        )
        if shoot_index is not None and shoot_index < len(shoot_rows):
            shoot_map = _zip_map(shoot_headers, shoot_rows[shoot_index])
        if basic_map.get("PTS") in (None, ""):
            continue
        values = [cell_value(basic_map.get(field, "")) for field in COUNTING_FIELDS]
        values += [cell_value(shoot_map.get(field, "")) for field in SHOOTING_FIELDS]
        rows.append([team, name] + values)
    return rows


def _status_label(raw):
    text = (raw or "").strip()
    mapping = {
        "OUT": "Out",
        "DAY-TO-DAY": "Day-to-day",
        "DTD": "Day-to-day",
        "QUESTIONABLE": "Questionable",
        "PROBABLE": "Probable",
    }
    return mapping.get(text.upper(), text)


def parse_injuries(html, team):
    soup = BeautifulSoup(html, "html.parser")
    rows = []
    current_date = ""
    for node in soup.select(".Injuries__groupDate, .ContentList__Item"):
        classes = node.get("class") or []
        if "Injuries__groupDate" in classes:
            current_date = node.get_text(" ", strip=True)
            continue
        name_node = node.select_one(".Athlete__PlayerName")
        status_node = node.select_one(".TextStatus")
        if not name_node or not status_node:
            continue
        player = name_node.get_text(strip=True)
        status = _status_label(status_node.get_text(strip=True))
        wrapper = node.select_one(".Athlete__PlayerWrapper")
        blob = wrapper.get_text(" ", strip=True) if wrapper else ""
        comment_match = re.search(r"Status\s+\S+\s*(.*)$", blob)
        comment = comment_match.group(1).strip() if comment_match else ""
        rows.append([team, player, comment, status, current_date])
    if rows:
        return rows

    for table in soup.find_all("table"):
        for tr in table.find_all("tr")[1:]:
            cells = tr.find_all(["td", "th"])
            if len(cells) < 4:
                continue
            player = cells[0].get_text(" ", strip=True)
            if not player or player.lower() in ("name", "player"):
                continue
            status = cells[1].get_text(" ", strip=True)
            date_val = cells[2].get_text(" ", strip=True)
            comment = cells[3].get_text(" ", strip=True)
            rows.append([team, player, comment, _status_label(status), date_val])
    return rows


def _selected_teams(argv):
    if "--only" in argv:
        index = argv.index("--only")
        if index + 1 >= len(argv):
            raise SystemExit("請在 --only 後面給隊名，例如 --only 賽爾提克")
        wanted = [name.strip() for name in argv[index + 1].split(",") if name.strip()]
        found = [team for team in NBA_TEAMS if team[0] in wanted]
        missing = [name for name in wanted if name not in {team[0] for team in found}]
        if missing:
            raise SystemExit("找不到球隊：%s" % "、".join(missing))
        return found
    return list(NBA_TEAMS)


def collect(teams, with_injuries):
    stat_rows = []
    injury_rows = []
    failures = []
    injury_pages_ok = 0
    for index, (team, abbr, slug) in enumerate(teams, start=1):
        print("[%02d/%02d] %s" % (index, len(teams), team))
        try:
            html = fetch_html(stats_url(abbr, slug), check_season=True)
            players = parse_player_stats(html, team)
        except Exception as exc:
            print("  球員狀態失敗：%s" % exc)
            failures.append((team, str(exc)))
            continue
        if not players:
            print("  球員狀態失敗：沒有球員列")
            failures.append((team, "沒有球員列"))
            continue
        stat_rows.extend(players)
        print("  球員 %d 人，例：%s PTS=%s" % (len(players), players[0][1], players[0][5]))
        if not with_injuries:
            continue
        try:
            injury_html = fetch_html(injuries_url(abbr, slug), check_season=False)
            injury_pages_ok += 1
            found = parse_injuries(injury_html, team)
            injury_rows.extend(found)
            print("  傷兵 %d 人" % len(found))
        except Exception as exc:
            print("  傷兵頁失敗（不覆寫傷兵）：%s" % exc)
            injury_pages_ok = -1
    return stat_rows, injury_rows, failures, injury_pages_ok


def _clear_below_header(sheet):
    if sheet.max_row >= 2:
        sheet.delete_rows(2, sheet.max_row - 1)


def write_workbook(path, stat_rows, injury_rows, write_injuries):
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    if "球員狀態" not in workbook.sheetnames:
        workbook.create_sheet("球員狀態")
    stats_sheet = workbook["球員狀態"]
    _clear_below_header(stats_sheet)
    for row_index, row in enumerate(stat_rows, start=2):
        for col_index, value in enumerate(row, start=1):
            stats_sheet.cell(row=row_index, column=col_index, value=value)

    if write_injuries:
        if "球員傷兵" not in workbook.sheetnames:
            workbook.create_sheet("球員傷兵")
        injury_sheet = workbook["球員傷兵"]
        injury_sheet["A1"] = "球隊"
        injury_sheet["B1"] = "球員"
        injury_sheet["C1"] = "傷勢"
        injury_sheet["D1"] = "狀態"
        injury_sheet["E1"] = "日期"
        _clear_below_header(injury_sheet)
        for row_index, row in enumerate(injury_rows, start=2):
            for col_index, value in enumerate(row[:5], start=1):
                injury_sheet.cell(row=row_index, column=col_index, value=value)

    workbook.save(path)
    workbook.close()


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    dry_run = "--dry-run" in argv
    with_injuries = "--stats-only" not in argv
    teams = _selected_teams(argv)
    if not dry_run and len(teams) != len(NBA_TEAMS):
        raise SystemExit("寫入 Excel 必須抓滿 30 隊。單隊請加 --dry-run。")

    print("=" * 60)
    print("爬取 ESPN 球員狀態" + ("與傷兵" if with_injuries else ""))
    print("鎖定 %s %s（season/%s seasontype/%s）" % (
        SEASON_LABEL, SEASON_TYPE_NAME, SEASON_YEAR, SEASON_TYPE,
    ))
    path = excel_path()
    print("Excel: %s" % path)
    print("=" * 60)

    stat_rows, injury_rows, failures, injury_pages_ok = collect(teams, with_injuries)
    if failures:
        print("未寫入 Excel。球員狀態失敗 %d 隊：" % len(failures))
        for team, reason in failures:
            print("  - %s：%s" % (team, reason))
        return 1

    write_injury_sheet = with_injuries and injury_pages_ok == len(teams)
    if dry_run:
        print("dry-run：球員狀態 %d 列，傷兵 %d 列，未寫入 Excel。" % (len(stat_rows), len(injury_rows)))
        return 0

    if not os.path.exists(path):
        print("錯誤：Excel 不存在。請先關閉並確認 OneDrive 檔在這條路徑。")
        return 1
    write_workbook(path, stat_rows, injury_rows, write_injury_sheet)
    print("完成：已覆蓋「球員狀態」%d 列。" % len(stat_rows))
    if write_injury_sheet:
        print("已覆蓋「球員傷兵」%d 列。" % len(injury_rows))
    elif with_injuries:
        print("傷兵頁沒有全部成功，保留原「球員傷兵」分頁。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
