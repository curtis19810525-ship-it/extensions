# -*- coding: utf-8 -*-
"""
爬取 ESPN 各隊 Team Total。
- 鎖定 2025-26 例行賽（season/2026、seasontype/2）
- 寫入分頁：各隊total
- 任一隊沒有 PTS 就不覆寫 Excel，避免再寫出「只有隊名、數字是空的」
"""

import os
import sys

from bs4 import BeautifulSoup

from espn_common import (
    NBA_TEAMS,
    SEASON_LABEL,
    SEASON_TYPE,
    SEASON_TYPE_NAME,
    SEASON_YEAR,
    cell_value,
    excel_path,
    fetch_html,
    stats_url,
)

STAT_FIELDS = [
    "GP", "GS", "MIN", "PTS", "OR", "REB", "AST", "STL", "BLK", "TO", "PF", "AST/TO",
    "FGM", "FGA", "FG%", "3PM", "3PA", "3P%", "FTM", "FTA", "FT%",
    "2PM", "2PA", "2P%", "SC-EFF", "SH-EFF",
]


def _text(cell):
    return " ".join(cell.get_text(" ", strip=True).split())


def _headers(table):
    return [_text(th).upper() for th in table.find_all("th")]


def _data_rows(table):
    body = table.find("tbody") or table
    rows = []
    for tr in body.find_all("tr"):
        cells = tr.find_all("td")
        if not cells:
            continue
        rows.append([_text(cell) for cell in cells])
    return rows


def _blocks(html):
    soup = BeautifulSoup(html, "html.parser")
    tables = soup.find_all("table")
    headers = [_headers(table) for table in tables]
    basic = shoot = None
    basic_need = {"GP", "PTS", "REB", "AST"}
    shoot_need = {"FGM", "FG%", "3PM", "FT%"}
    for index in range(len(tables) - 1):
        left, right = headers[index], headers[index + 1]
        if left == ["NAME"] and basic_need.issubset(set(right)) and basic is None:
            basic = (_data_rows(tables[index]), right, _data_rows(tables[index + 1]))
        if left == ["NAME"] and shoot_need.issubset(set(right)) and shoot is None:
            shoot = (_data_rows(tables[index]), right, _data_rows(tables[index + 1]))
    return basic, shoot


def _map_at(names, headers, rows, label):
    for index, name_cells in enumerate(names):
        if not name_cells or name_cells[0] != label:
            continue
        if index >= len(rows):
            return {}
        values = rows[index]
        width = min(len(headers), len(values))
        return {headers[i]: values[i] for i in range(width)}
    return {}


def parse_team_total(html):
    basic, shoot = _blocks(html)
    if not basic or not shoot:
        return {}
    basic_map = _map_at(basic[0], basic[1], basic[2], "Total")
    shoot_map = _map_at(shoot[0], shoot[1], shoot[2], "Total")
    merged = {}
    merged.update(basic_map)
    merged.update(shoot_map)
    return merged


def build_rows(teams):
    rows = []
    failures = []
    for index, (team, abbr, slug) in enumerate(teams, start=1):
        url = stats_url(abbr, slug)
        print("[%02d/%02d] %s" % (index, len(teams), team))
        try:
            html = fetch_html(url, check_season=True)
            stats = parse_team_total(html)
        except Exception as exc:
            print("  失敗：%s" % exc)
            failures.append((team, str(exc)))
            continue
        if stats.get("PTS") in (None, "") or stats.get("GP") in (None, ""):
            print("  失敗：沒有 Total 的 GP/PTS")
            failures.append((team, "沒有 Total 的 GP/PTS"))
            continue
        row = [team, "Total"] + [cell_value(stats.get(field, "")) for field in STAT_FIELDS]
        rows.append(row)
        print("  GP=%s PTS=%s" % (row[2], row[5]))
    return rows, failures


def write_rows(path, rows):
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    if "各隊total" not in workbook.sheetnames:
        workbook.create_sheet("各隊total")
    sheet = workbook["各隊total"]
    if sheet.max_row >= 2:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=col_index, value=value)
    workbook.save(path)
    workbook.close()


def _selected_teams(argv):
    if "--only" in argv:
        index = argv.index("--only")
        if index + 1 >= len(argv):
            raise SystemExit("請在 --only 後面給隊名，例如 --only 賽爾提克")
        wanted = argv[index + 1].split(",")
        wanted = [name.strip() for name in wanted if name.strip()]
        found = [team for team in NBA_TEAMS if team[0] in wanted]
        missing = [name for name in wanted if name not in {team[0] for team in found}]
        if missing:
            raise SystemExit("找不到球隊：%s" % "、".join(missing))
        return found
    return list(NBA_TEAMS)


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    dry_run = "--dry-run" in argv
    teams = _selected_teams(argv)
    if not dry_run and len(teams) != len(NBA_TEAMS):
        raise SystemExit("寫入 Excel 必須抓滿 30 隊。單隊請加 --dry-run。")

    print("=" * 60)
    print("爬取 ESPN 各隊 Total")
    print("鎖定 %s %s（season/%s seasontype/%s）" % (
        SEASON_LABEL, SEASON_TYPE_NAME, SEASON_YEAR, SEASON_TYPE,
    ))
    path = excel_path()
    print("Excel: %s" % path)
    print("=" * 60)

    rows, failures = build_rows(teams)
    if failures or len(rows) != len(teams):
        print("未寫入 Excel。失敗 %d 隊：" % len(failures))
        for team, reason in failures:
            print("  - %s：%s" % (team, reason))
        return 1

    if dry_run:
        print("dry-run：%d 隊有數字，未寫入 Excel。" % len(rows))
        return 0

    if not os.path.exists(path):
        print("錯誤：Excel 不存在。請先關閉並確認 OneDrive 檔在這條路徑。")
        return 1
    write_rows(path, rows)
    print("完成：已覆蓋寫入分頁「各隊total」%d 隊。" % len(rows))
    return 0


if __name__ == "__main__":
    sys.exit(main())
